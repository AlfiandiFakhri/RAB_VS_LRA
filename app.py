import streamlit as st
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import io
import re
from copy import copy
import plotly.express as px
import plotly.graph_objects as go
import numpy as np
import matplotlib.pyplot as plt
from difflib import SequenceMatcher

# Pengaman untuk Library PowerPoint
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

# ==========================================
# KONFIGURASI HALAMAN
# ==========================================
st.set_page_config(page_title="RAB vs LRA", layout="wide", page_icon="📊")

def normalize_text(text):
    t = str(text)
    t = re.sub(r'^[\w\.]+\s*-\s*', '', t)
    t = re.sub(r'^[a-zA-Z0-9]+\.\s*', '', t)
    t = t.split('|')[0]
    t = t.replace('-', '').strip().lower().replace(',', '')
    return t

def safe_float(val):
    if pd.isna(val): return 0.0
    try:
        if isinstance(val, str):
            val = val.replace(',', '').strip()
            if val == '-' or val == '': return 0.0
        return float(val)
    except (ValueError, TypeError):
        return 0.0

def match_texts_smart(t1, t2):
    if not t1 or not t2: return False
    clean_t1 = normalize_text(t1)
    clean_t2 = normalize_text(t2)
    
    if len(clean_t1) < 3 or len(clean_t2) < 3:
        return clean_t1 == clean_t2
    
    if clean_t1 == clean_t2 or clean_t1 in clean_t2 or clean_t2 in clean_t1: return True
        
    similarity = SequenceMatcher(None, clean_t1, clean_t2).ratio()
    if similarity >= 0.65: return True
        
    w1 = set(clean_t1.split())
    w2 = set(clean_t2.split())
    if len(w1) > 0 and len(w2) > 0:
        shorter = w1 if len(w1) < len(w2) else w2
        longer = w2 if len(w1) < len(w2) else w1
        overlap = len(shorter.intersection(longer))
        if (overlap / len(shorter)) >= 0.70: return True
    return False

def get_row_pagu(ws, row_idx):
    try:
        f_val = ws.cell(row=row_idx, column=6).value
        h_val = ws.cell(row=row_idx, column=8).value
        r_val = ws.cell(row=row_idx, column=18).value
        
        def parse_val(v):
            if isinstance(v, (int, float)): return float(v)
            if isinstance(v, str) and v.startswith('='):
                parts = v.lstrip('=').split('*')
                nums = []
                for p in parts:
                    p = p.strip()
                    try:
                        nums.append(float(p))
                    except:
                        col_letter = ''.join([c for c in p if c.isalpha()])
                        row_num = int(''.join([c for c in p if c.isdigit()]))
                        col_idx = ord(col_letter.upper()) - 64
                        cell_v = ws.cell(row=row_num, column=col_idx).value
                        if isinstance(cell_v, (int, float)): nums.append(float(cell_v))
                        else: nums.append(1.0)
                res = 1.0
                for n in nums: res *= n
                return res
            return 0.0
            
        vol1 = parse_val(f_val) if f_val is not None else 1.0
        vol2 = parse_val(h_val) if h_val is not None else 1.0
        biaya = parse_val(r_val) if r_val is not None else 0.0
        
        h_raw = ws.cell(row=row_idx, column=7).value
        if h_raw is not None and isinstance(h_raw, str) and 'x' in h_raw.lower():
            return vol1 * vol2 * biaya
        else:
            return vol1 * biaya
    except:
        return 0.0

@st.cache_data
def parse_lra_files_cached(lra_files_input, list_semua_bulan):
    data_realisasi = {}
    satker_summary = {"pagu": 0, "realisasi": 0, "sisa": 0, "outstanding": 0}
    monthly_totals = {b: 0 for b in list_semua_bulan}
    component_summary = {}
    component_metrics = {}
    sub_component_realisasi = {}
    sub_component_sisa = {} 
    
    for fname, fbytes in lra_files_input:
        file_obj = io.BytesIO(fbytes)
        df_raw = pd.read_excel(file_obj, header=None, nrows=15)
        file_obj.seek(0)
        
        bulan_file = None
        is_outstanding_file = False
        
        fname_upper = fname.upper()
        for b in list_semua_bulan:
            if b in fname_upper:
                bulan_file = b
                break

        if not bulan_file:
            for r in range(len(df_raw)):
                row_text_raw = " ".join(str(val) for val in df_raw.iloc[r].values if pd.notna(val)).upper()
                for b in list_semua_bulan:
                    if b in row_text_raw:
                        bulan_file = b
                        break
                if bulan_file: break

        fname_clean = fname_upper.replace(" ", "")
        keyword_out = ["SEMUA", "LEVEL", "ALL", "PERIODE", "REKAP", "GUP"]
        if any(kw in fname_clean for kw in keyword_out) and not bulan_file:
            is_outstanding_file = True

        for r in range(len(df_raw)):
            row_text_raw = " ".join(str(val) for val in df_raw.iloc[r].values if pd.notna(val)).upper()
            if any(k in row_text_raw for k in ["SEMUA", "LEVEL", "ALL", "PERIODE"]):
                if not bulan_file:
                    is_outstanding_file = True
                    break

        df_lra = pd.read_excel(file_obj, skiprows=5)
        
        kolom_lra = [str(col).upper() for col in df_lra.columns]
        ada_gup_spm = any('GUP' in col or 'SPM' in col or 'VERIFIKASI' in col for col in kolom_lra)
        if ada_gup_spm and not bulan_file:
            is_outstanding_file = True

        if 'Level' in df_lra.columns:
            current_komp = ""
            for _, row_lra in df_lra.iterrows():
                lvl = str(row_lra.get('Level')).strip()
                uraian = str(row_lra.get('Kode / Uraian', '')).strip()
                
                pagu = safe_float(row_lra.get('Pagu'))
                realisasi_sub = safe_float(row_lra.get('Total Realisasi'))
                sisa_sub = safe_float(row_lra.get('Sisa'))
                
                if lvl == 'Komponen':
                    current_komp = uraian
                    component_summary[current_komp] = pagu
                    if is_outstanding_file or current_komp not in component_metrics or realisasi_sub > component_metrics[current_komp]['realisasi']:
                        component_metrics[current_komp] = {
                            "pagu": pagu,
                            "realisasi": realisasi_sub,
                            "sisa": sisa_sub,
                            "persen": (realisasi_sub / pagu * 100) if pagu > 0 else 0
                        }
                elif lvl == 'Sub Komponen':
                    if current_komp:
                        if current_komp not in sub_component_realisasi: sub_component_realisasi[current_komp] = {}
                        if current_komp not in sub_component_sisa: sub_component_sisa[current_komp] = {}
                            
                        current_stored_real = sub_component_realisasi[current_komp].get(uraian, -1)
                        if is_outstanding_file or uraian not in sub_component_realisasi[current_komp] or realisasi_sub > current_stored_real:
                            sub_component_realisasi[current_komp][uraian] = realisasi_sub
                            sub_component_sisa[current_komp][uraian] = sisa_sub

        if is_outstanding_file:
            satker_row = df_lra[df_lra['Level'].astype(str).str.strip() == 'Satker']
            if not satker_row.empty:
                satker_summary["pagu"] = safe_float(satker_row['Pagu'].values[0])
                satker_summary["realisasi"] = safe_float(satker_row['Total Realisasi'].values[0])
                satker_summary["sisa"] = safe_float(satker_row['Sisa'].values[0])
            
            detail_rows = df_lra[df_lra['Level'].astype(str).str.strip() == 'Detail']
            satker_summary["outstanding"] = (
                pd.to_numeric(detail_rows['GUP'], errors='coerce').fillna(0).sum() +
                pd.to_numeric(detail_rows['SPM'], errors='coerce').fillna(0).sum() +
                pd.to_numeric(detail_rows['Verifikasi'], errors='coerce').fillna(0).sum()
            )

        if not bulan_file and not is_outstanding_file: continue
            
        cur_komp, cur_sub, cur_akun = "GLOBAL", "GLOBAL", "GLOBAL"
        
        for index, row in df_lra.iterrows():
            lvl = str(row.get('Level')).strip()
            uraian = str(row.get('Kode / Uraian', '')).strip()
            
            realisasi = safe_float(row.get('Total Realisasi'))
            outstanding_val = safe_float(row.get('GUP')) + safe_float(row.get('SPM')) + safe_float(row.get('Verifikasi'))
                
            if lvl == 'Komponen':
                match = re.search(r'(\d{3})\s*-', uraian)
                if match: cur_komp = match.group(1)
                cur_sub, cur_akun = "GLOBAL", "GLOBAL"
            elif lvl == 'Sub Komponen':
                match = re.search(r'([A-Z])\s*-', uraian)
                if match: cur_sub = match.group(1)
                cur_akun = "GLOBAL"
            elif lvl == 'Akun':
                match = re.search(r'(\d{6})\s*-', uraian)
                if match: cur_akun = match.group(1)
            
            kamar_unik = f"{cur_komp}_{cur_sub}_{cur_akun}"
            if kamar_unik not in data_realisasi: data_realisasi[kamar_unik] = {}
                
            if pd.notna(uraian) and uraian not in ['nan', '']:
                norm_lra = normalize_text(uraian)
                
                if norm_lra not in data_realisasi[kamar_unik]:
                    data_realisasi[kamar_unik][norm_lra] = {b: 0 for b in list_semua_bulan}
                    data_realisasi[kamar_unik][norm_lra]['OUTSTANDING'] = 0
                
                if bulan_file:
                    data_realisasi[kamar_unik][norm_lra][bulan_file] += realisasi
                    monthly_totals[bulan_file] += realisasi
                if is_outstanding_file:
                    data_realisasi[kamar_unik][norm_lra]['OUTSTANDING'] += outstanding_val
                
    return data_realisasi, satker_summary, monthly_totals, component_summary, component_metrics, sub_component_realisasi, sub_component_sisa

@st.cache_data
def process_rab_lra_cached(rab_bytes, data_realisasi_tuple, list_semua_bulan):
    file_rab_io = io.BytesIO(rab_bytes)
    wb = load_workbook(file_rab_io)
    ws = wb.active 
    
    baris_header = 13 
    baris_mulai_data = 14
    
    kolom_baru = ["TOTAL Realisasi", "SISA", "OUT STANDING", "JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI", "JULI", "AGUSTUS", "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER", "KETERANGAN"]
    header_ref = ws.cell(row=baris_header, column=19)
    start_col = 20
    
    warna_hijau = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
    warna_kuning = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    warna_orange = PatternFill(start_color="FCD5B4", end_color="FCD5B4", fill_type="solid") 
    
    for i, nama_kolom in enumerate(kolom_baru):
        cell = ws.cell(row=baris_header, column=start_col + i)
        cell.value = nama_kolom
        if header_ref.has_style:
            cell.font = copy(header_ref.font)
            cell.border = copy(header_ref.border)
            cell.alignment = copy(header_ref.alignment)
        if nama_kolom == "TOTAL Realisasi": cell.fill = warna_hijau
        elif nama_kolom == "SISA": cell.fill = warna_kuning
        elif nama_kolom == "OUT STANDING": cell.fill = warna_orange
        else:
            if header_ref.has_style: cell.fill = copy(header_ref.fill)

    ws.column_dimensions[get_column_letter(20)].width = 20.5 
    ws.column_dimensions[get_column_letter(21)].width = 18.2 
    ws.column_dimensions[get_column_letter(22)].width = 18.2 
    for c_idx in range(23, 35): ws.column_dimensions[get_column_letter(c_idx)].width = 15.0 
    ws.column_dimensions[get_column_letter(35)].width = 35.0 

    max_row = ws.max_row
    data_ditemukan = 0
    rab_komp, rab_sub, rab_akun = "GLOBAL", "GLOBAL", "GLOBAL"
    baris_terpakai = set()
    summary_preview = []
    
    import copy as cp
    data_realisasi = cp.deepcopy(data_realisasi_tuple)

    for row_idx in range(baris_mulai_data, max_row + 1):
        kode_col = str(ws.cell(row=row_idx, column=2).value).strip()
        if kode_col.isdigit() and len(kode_col) == 3: 
            rab_komp = kode_col
            rab_sub, rab_akun = "GLOBAL", "GLOBAL"
        elif kode_col.isdigit() and len(kode_col) == 6: rab_akun = kode_col

        sub_col = str(ws.cell(row=row_idx, column=4).value).strip()
        if re.match(r'^[A-Z]\.?\s*$', sub_col):
            rab_sub = sub_col.replace('.', '').strip()
            rab_akun = "GLOBAL" 
            
        kamar_rab_saat_ini = f"{rab_komp}_{rab_sub}_{rab_akun}"
        
        bagian_teks = []
        for col_idx in range(3, 7):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val and isinstance(val, str) and str(val).strip() not in ['-', '']: bagian_teks.append(str(val).strip())
        
        if bagian_teks:
            uraian_rab = " ".join(bagian_teks)
            matched_key = None
            kamar_opsi = [kamar_rab_saat_ini, f"{rab_komp}_{rab_sub}_GLOBAL", f"{rab_komp}_GLOBAL_GLOBAL"]
            
            for kamar in kamar_opsi:
                if kamar in data_realisasi:
                    for key, dict_bulanan in data_realisasi[kamar].items():
                        if match_texts_smart(key, uraian_rab):
                            matched_key = (kamar, key)
                            break
                if matched_key: break
            
            if matched_key:
                kamar_ketemu, key_ketemu = matched_key
                dict_bulanan = data_realisasi[kamar_ketemu].pop(key_ketemu) 
                nilai_outstanding = dict_bulanan.pop('OUTSTANDING', 0)
                
                target_row = row_idx
                for r_cek in range(row_idx, min(row_idx + 6, max_row + 1)):
                    val_s = ws.cell(row=r_cek, column=19).value
                    if val_s is not None and r_cek not in baris_terpakai:
                        target_row = r_cek
                        break
                baris_terpakai.add(target_row)
                
                ws.cell(row=target_row, column=22).value = nilai_outstanding
                row_bulanan_val = {}
                for idx_b, b_name in enumerate(list_semua_bulan):
                    val_b = dict_bulanan[b_name]
                    ws.cell(row=target_row, column=23 + idx_b).value = val_b
                    row_bulanan_val[b_name] = val_b
                
                ws.cell(row=target_row, column=20).value = f"=SUM(W{target_row}:AH{target_row})+V{target_row}"
                ws.cell(row=target_row, column=21).value = f"=S{target_row}-T{target_row}"

                pagu_val = get_row_pagu(ws, target_row)
                summary_preview.append({
                    "Komponen": rab_komp,
                    "Uraian": uraian_rab,
                    "Pagu": pagu_val,
                    "Realisasi": sum(row_bulanan_val.values()) + nilai_outstanding,
                    "Outstanding": nilai_outstanding,
                    **row_bulanan_val
                })

                data_ref = ws.cell(row=target_row, column=19)
                for c_idx in range(20, 36):
                    c = ws.cell(row=target_row, column=c_idx)
                    if data_ref.has_style:
                        c.font, c.border, c.alignment, c.number_format = copy(data_ref.font), copy(data_ref.border), copy(data_ref.alignment), copy(data_ref.number_format)
                data_ditemukan += 1

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue(), data_ditemukan, pd.DataFrame(summary_preview)

# ==========================================
# FUNGSI PEMBUATAN FILE POWERPOINT (PPTX) 
# ==========================================
def create_powerpoint_report(satker_summary, component_metrics, monthly_totals, sub_component_realisasi, sub_component_sisa):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    # --- WARNA TEMA DEEP TEAL ---
    DEEP_TEAL = RGBColor(0, 64, 64)       
    GOLD_ACCENT = RGBColor(218, 165, 32)  
    GREY_TEXT = RGBColor(80, 90, 100)     
    WHITE = RGBColor(255, 255, 255)       
    LIGHT_BG = RGBColor(245, 247, 250)    

    def add_slide_header(slide, title_text):
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.8))
        p = t_box.text_frame.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.name = "Arial"
        p.font.color.rgb = DEEP_TEAL
        
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = GOLD_ACCENT
        line.line.fill.background()
    
    # ---------------------------------------------------------
    # 1. SLIDE COVER (JUDUL)
    # ---------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid(); bg.fill.fore_color.rgb = DEEP_TEAL; bg.line.fill.background()
    
    accent = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(4.2), Inches(1.5), Inches(0.08))
    accent.fill.solid(); accent.fill.fore_color.rgb = GOLD_ACCENT; accent.line.fill.background()

    tf1 = slide1.shapes.add_textbox(Inches(0.9), Inches(2.2), Inches(11), Inches(2)).text_frame
    tf1.word_wrap = True
    
    p1 = tf1.paragraphs[0]
    p1.text = "LAPORAN EKSEKUTIF"
    p1.font.size, p1.font.bold, p1.font.name, p1.font.color.rgb = Pt(48), True, "Arial", WHITE
    
    p2 = tf1.add_paragraph()
    p2.text = "Konsolidasi Anggaran (RAB vs LRA)"
    p2.font.size, p2.font.name, p2.font.color.rgb = Pt(24), "Arial", RGBColor(180, 205, 205)
    
    # ---------------------------------------------------------
    # 2. SLIDE RINGKASAN EKSEKUTIF 
    # ---------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide2, "Ringkasan Kinerja Anggaran")
    
    total_pagu = satker_summary["pagu"]
    total_realisasi = satker_summary["realisasi"]
    total_sisa = satker_summary["sisa"]
    persen_nasional = (total_realisasi / total_pagu * 100) if total_pagu > 0 else 0
    
    metrics = [
        ("Total Pagu Anggaran", f"Rp {total_pagu:,.0f}"),
        ("Total Realisasi", f"Rp {total_realisasi:,.0f}"),
        ("Sisa Anggaran", f"Rp {total_sisa:,.0f}"),
        ("Tingkat Penyerapan", f"{persen_nasional:.2f}%")
    ]
    
    for i, (label, value) in enumerate(metrics):
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + (i * 3.0)), Inches(1.5), Inches(2.8), Inches(1.1))
        card.fill.solid(); card.fill.fore_color.rgb = LIGHT_BG; card.line.color.rgb = RGBColor(220, 225, 230)
        c_tf = card.text_frame; c_tf.clear()
        
        c_p1 = c_tf.paragraphs[0]
        c_p1.text, c_p1.font.size, c_p1.font.color.rgb = label, Pt(12), GREY_TEXT
        
        c_p2 = c_tf.add_paragraph()
        c_p2.text, c_p2.font.size, c_p2.font.bold, c_p2.font.color.rgb = value, Pt(18), True, DEEP_TEAL
    
    if component_metrics:
        p_c = slide2.shapes.add_textbox(Inches(0.8), Inches(2.9), Inches(11.7), Inches(0.5)).text_frame.paragraphs[0]
        p_c.text, p_c.font.size, p_c.font.bold, p_c.font.color.rgb = "Rincian per Komponen Utama:", Pt(16), True, DEEP_TEAL
        
        rows = len(component_metrics) + 1
        cols = 5
        c_table = slide2.shapes.add_table(rows, cols, Inches(0.8), Inches(3.4), Inches(11.7), Inches(0.4 * rows)).table
        c_table.columns[0].width, c_table.columns[1].width, c_table.columns[2].width, c_table.columns[3].width, c_table.columns[4].width = Inches(4.5), Inches(2.0), Inches(2.0), Inches(2.0), Inches(1.2)
        
        headers = ["Komponen Utama", "Pagu Anggaran (Rp)", "Total Realisasi (Rp)", "Sisa Anggaran (Rp)", "% Realisasi"]
        for col_idx, h_text in enumerate(headers):
            cell = c_table.cell(0, col_idx); cell.text = h_text
            cell.fill.solid(); cell.fill.fore_color.rgb = DEEP_TEAL
            p = cell.text_frame.paragraphs[0]
            p.font.bold, p.font.color.rgb, p.font.size = True, WHITE, Pt(12)
            
        for r_idx, (komp_name, m) in enumerate(component_metrics.items(), start=1):
            row_color = WHITE if r_idx % 2 != 0 else LIGHT_BG
            data = [komp_name, f"Rp {m['pagu']:,.0f}", f"Rp {m['realisasi']:,.0f}", f"Rp {m['sisa']:,.0f}", f"{m['persen']:.2f}%"]
            for c_idx, val in enumerate(data):
                cell = c_table.cell(r_idx, c_idx); cell.text = val
                cell.fill.solid(); cell.fill.fore_color.rgb = row_color
                p = cell.text_frame.paragraphs[0]
                p.font.color.rgb, p.font.size = GREY_TEXT, Pt(11)

    # ---------------------------------------------------------
    # 3. SLIDE GRAFIK TREN BULANAN
    # ---------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_slide_header(slide3, "Tren Penyerapan Anggaran Bulanan")
    
    plt.rcParams['font.family'] = 'sans-serif'
    fig_m, ax_m = plt.subplots(figsize=(12, 5))
    
    months = list(monthly_totals.keys())
    values = list(monthly_totals.values())
    ax_m.bar(months, [v / 1e9 for v in values], color='#008080', width=0.6) 
    
    ax_m.spines['top'].set_visible(False)
    ax_m.spines['right'].set_visible(False)
    ax_m.spines['left'].set_color('#DDDDDD')
    ax_m.spines['bottom'].set_color('#DDDDDD')
    ax_m.tick_params(bottom=False, left=False)
    
    ax_m.set_ylabel('Realisasi (Miliar Rp)', fontsize=12, fontweight='bold', color='#333333')
    plt.xticks(rotation=30, ha='right', fontsize=11, color='#555555')
    plt.yticks(fontsize=11, color='#555555')
    ax_m.grid(axis='y', linestyle='--', alpha=0.4)
    
    plt.tight_layout()
    img_buf = io.BytesIO(); fig_m.savefig(img_buf, format='png', dpi=300, bbox_inches='tight', transparent=True); plt.close(fig_m); img_buf.seek(0)
    slide3.shapes.add_picture(img_buf, Inches(0.5), Inches(1.5), width=Inches(12.0))
    
    # ---------------------------------------------------------
    # 4. SLIDE RINCIAN SUB KOMPONEN (REALISASI & SISA)
    # ---------------------------------------------------------
    semua_komponen_unik = sorted(list(set(list(sub_component_realisasi.keys()) + list(sub_component_sisa.keys()))))
    
    for komp_name in semua_komponen_unik:
        
        # ---- SLIDE: REALISASI SUB KOMPONEN ----
        if komp_name in sub_component_realisasi:
            sub_dict_real_table = {k: float(v) for k, v in sub_component_realisasi[komp_name].items() if pd.notna(v) and float(v) >= 0}
            
            if sub_dict_real_table:
                slide_sub = prs.slides.add_slide(blank_layout)
                add_slide_header(slide_sub, f"Realisasi Sub Komponen: {komp_name}")
                
                sub_labels_chart = [k for k, v in sub_dict_real_table.items() if v > 0]
                sub_vals_chart = [v for v in sub_dict_real_table.values() if v > 0]
                total_komp_chart = sum(sub_vals_chart)
                
                if total_komp_chart > 0:
                    fig_p, ax_p = plt.subplots(figsize=(6, 5))
                    colors_real = ['#004040', '#006666', '#008080', '#20B2AA', '#48D1CC', '#5F9EA0', '#8FBC8F']
                    
                    labels_with_pct = [f"Sub {l.split('-')[0].strip()}\n({val/total_komp_chart*100:.1f}%)" for l, val in zip(sub_labels_chart, sub_vals_chart)]
                    
                    wedges, texts = ax_p.pie(
                        sub_vals_chart, 
                        labels=labels_with_pct, 
                        startangle=90, 
                        colors=colors_real, 
                        wedgeprops=dict(width=0.4, edgecolor='w'),
                        labeldistance=1.05
                    )
                    
                    plt.setp(texts, size=10, weight="bold", color="#333333")
                    ax_p.axis('equal'); plt.tight_layout()
                    
                    pie_buf = io.BytesIO(); fig_p.savefig(pie_buf, format='png', dpi=300, bbox_inches='tight', transparent=True); plt.close(fig_p); pie_buf.seek(0)
                    slide_sub.shapes.add_picture(pie_buf, Inches(0.5), Inches(1.8), width=Inches(5.0))
                else:
                    t_box_empty = slide_sub.shapes.add_textbox(Inches(1.5), Inches(3.0), Inches(4.0), Inches(1.0)).text_frame
                    t_box_empty.text = "Seluruh sub komponen belum memiliki realisasi (Rp 0)"
                
                sub_labels_table = list(sub_dict_real_table.keys())
                sub_vals_table = list(sub_dict_real_table.values())
                total_komp_table = sum(sub_vals_table)
                
                rows = len(sub_dict_real_table) + 1
                table = slide_sub.shapes.add_table(rows, 3, Inches(5.8), Inches(1.8), Inches(7.0), Inches(0.4 * rows)).table
                table.columns[0].width, table.columns[1].width, table.columns[2].width = Inches(3.5), Inches(2.0), Inches(1.5)
                
                for col_idx, h_text in enumerate(["Sub Komponen", "Total Realisasi (Rp)", "Persentase"]):
                    cell = table.cell(0, col_idx); cell.text = h_text
                    cell.fill.solid(); cell.fill.fore_color.rgb = DEEP_TEAL
                    p = cell.text_frame.paragraphs[0]
                    p.font.bold, p.font.color.rgb, p.font.size = True, WHITE, Pt(12)
                
                for r_idx, (lbl, val) in enumerate(zip(sub_labels_table, sub_vals_table), start=1):
                    pct = (val / total_komp_table * 100) if total_komp_table > 0 else 0
                    row_color = WHITE if r_idx % 2 != 0 else LIGHT_BG
                    for c_idx, text_val in enumerate([lbl, f"Rp {val:,.0f}", f"{pct:.2f}%"]):
                        cell = table.cell(r_idx, c_idx); cell.text = text_val
                        cell.fill.solid(); cell.fill.fore_color.rgb = row_color
                        p = cell.text_frame.paragraphs[0]
                        p.font.color.rgb, p.font.size = GREY_TEXT, Pt(11)

        # ---- SLIDE: SISA ANGGARAN SUB KOMPONEN ----
        if komp_name in sub_component_sisa:
            sub_dict_sisa_table = {k: float(v) for k, v in sub_component_sisa[komp_name].items() if pd.notna(v) and float(v) >= 0}
            
            if sub_dict_sisa_table:
                slide_sisa = prs.slides.add_slide(blank_layout)
                add_slide_header(slide_sisa, f"Sisa Anggaran Sub Komponen: {komp_name}")
                
                sub_labels_s_chart = [k for k, v in sub_dict_sisa_table.items() if v > 0]
                sub_vals_s_chart = [v for v in sub_dict_sisa_table.values() if v > 0]
                total_komp_s_chart = sum(sub_vals_s_chart)
                
                if total_komp_s_chart > 0:
                    fig_s, ax_s = plt.subplots(figsize=(6, 5))
                    colors_sisa = ['#DAA520', '#CD853F', '#D2691E', '#B8860B', '#8B4513', '#A0522D', '#D2B48C']
                    
                    labels_s_with_pct = [f"Sub {l.split('-')[0].strip()}\n({val/total_komp_s_chart*100:.1f}%)" for l, val in zip(sub_labels_s_chart, sub_vals_s_chart)]
                    
                    wedges_s, texts_s = ax_s.pie(
                        sub_vals_s_chart, 
                        labels=labels_s_with_pct, 
                        startangle=90, 
                        colors=colors_sisa, 
                        wedgeprops=dict(width=0.4, edgecolor='w'),
                        labeldistance=1.05
                    )
                    
                    plt.setp(texts_s, size=10, weight="bold", color="#333333")
                    ax_s.axis('equal'); plt.tight_layout()
                    
                    pie_buf_s = io.BytesIO(); fig_s.savefig(pie_buf_s, format='png', dpi=300, bbox_inches='tight', transparent=True); plt.close(fig_s); pie_buf_s.seek(0)
                    slide_sisa.shapes.add_picture(pie_buf_s, Inches(0.5), Inches(1.8), width=Inches(5.0))
                else:
                    t_box_empty_s = slide_sisa.shapes.add_textbox(Inches(1.5), Inches(3.0), Inches(4.0), Inches(1.0)).text_frame
                    t_box_empty_s.text = "Seluruh anggaran telah terealisasi (Sisa Rp 0)"
                
                sub_labels_s_table = list(sub_dict_sisa_table.keys())
                sub_vals_s_table = list(sub_dict_sisa_table.values())
                total_komp_s_table = sum(sub_vals_s_table)
                
                rows_s = len(sub_dict_sisa_table) + 1
                table_s = slide_sisa.shapes.add_table(rows_s, 3, Inches(5.8), Inches(1.8), Inches(7.0), Inches(0.4 * rows_s)).table
                table_s.columns[0].width, table_s.columns[1].width, table_s.columns[2].width = Inches(3.5), Inches(2.0), Inches(1.5)
                
                for col_idx, h_text in enumerate(["Sub Komponen", "Total Sisa (Rp)", "Persentase Sisa"]):
                    cell = table_s.cell(0, col_idx); cell.text = h_text
                    cell.fill.solid(); cell.fill.fore_color.rgb = GOLD_ACCENT 
                    p = cell.text_frame.paragraphs[0]
                    p.font.bold, p.font.color.rgb, p.font.size = True, WHITE, Pt(12)
                
                for r_idx, (lbl, val) in enumerate(zip(sub_labels_s_table, sub_vals_s_table), start=1):
                    pct = (val / total_komp_s_table * 100) if total_komp_s_table > 0 else 0
                    row_color = WHITE if r_idx % 2 != 0 else LIGHT_BG
                    for c_idx, text_val in enumerate([lbl, f"Rp {val:,.0f}", f"{pct:.2f}%"]):
                        cell = table_s.cell(r_idx, c_idx); cell.text = text_val
                        cell.fill.solid(); cell.fill.fore_color.rgb = row_color
                        p = cell.text_frame.paragraphs[0]
                        p.font.color.rgb, p.font.size = GREY_TEXT, Pt(11)
                
    ppt_output = io.BytesIO(); prs.save(ppt_output); ppt_output.seek(0)
    return ppt_output.getvalue()

# ==========================================
# ANTARMUKA PENGGUNA (UI STREAMLIT)
# ==========================================
st.title("📊 LAPORAN REALISASI")
st.markdown("**Executive Summary**")
st.divider()

col1, col2 = st.columns(2)
with col1:
    st.info("File RAB")
    file_rab = st.file_uploader("Upload Excel RAB", type=['xlsx', 'xls'], key="rab")
with col2:
    st.info("File LRA")
    file_lra_list = st.file_uploader("Upload Excel LRA (Bisa pilih banyak file)", type=['xlsx', 'xls'], accept_multiple_files=True, key="lra")

if file_rab and file_lra_list:
    st.divider()
    if st.button("🚀 PROSES", type="primary", use_container_width=True):
        
        list_semua_bulan = ['JANUARI', 'FEBRUARI', 'MARET', 'APRIL', 'MEI', 'JUNI', 'JULI', 'AGUSTUS', 'SEPTEMBER', 'OKTOBER', 'NOVEMBER', 'DESEMBER']
        
        try:
            with st.status("Sedang memproses dokumen dan menyusun ringkasan (Cached)...", expanded=True) as status:
                st.write("Mengekstrak data dari seluruh LRA (Realisasi & Outstanding)...")
                lra_files_tuple = tuple((f.name, f.getvalue()) for f in file_lra_list)
                list_bulan_tuple = tuple(list_semua_bulan)
                
                data_realisasi, satker_summary, monthly_totals, component_summary, component_metrics, sub_component_realisasi, sub_component_sisa = parse_lra_files_cached(lra_files_tuple, list_bulan_tuple)
                
                st.write("Menyelaraskan ...")
                rab_bytes = file_rab.getvalue()
                output_excel_bytes, data_ditemukan, df_preview = process_rab_lra_cached(rab_bytes, data_realisasi, list_bulan_tuple)
                
                st.write("Menyiapkan dokumen presentasi PowerPoint (.pptx)...")
                ppt_output_bytes = create_powerpoint_report(satker_summary, component_metrics, monthly_totals, sub_component_realisasi, sub_component_sisa)
                
                status.update(label="Proses Selesai!", state="complete", expanded=False)

            st.success(f"🎉 SUKSES!")
            st.divider()

            # ==========================================
            # DASHBOARD EXECUTIVE SUMMARY METRICS
            # ==========================================
            st.subheader("📈 Executive Summary")
            total_pagu_all = satker_summary["pagu"]; total_realisasi_incl_out = satker_summary["realisasi"]; total_sisa_all = satker_summary["sisa"]
            persen_nasional = (total_realisasi_incl_out / total_pagu_all * 100) if total_pagu_all > 0 else 0

            m1, m2, m3, m4 = st.columns(4)
            m1.metric("💰 Total Pagu Anggaran", f"Rp {total_pagu_all:,.0f}")
            m2.metric("📉 Total Realisasi", f"Rp {total_realisasi_incl_out:,.0f}")
            m3.metric("🟡 Sisa Anggaran", f"Rp {total_sisa_all:,.0f}")
            m4.metric("📊 Persentase", f"{persen_nasional:.2f}%")

            if component_metrics:
                st.markdown("---")
                st.markdown("### 🏷️ Summary Berdasarkan Kegiatan")
                comp_cols = st.columns(2)
                for idx, (komp_name, metrics) in enumerate(component_metrics.items()):
                    with comp_cols[idx % 2]:
                        st.markdown(f"**{komp_name}**")
                        sub_c1, sub_c2 = st.columns(2)
                        sub_c1.metric("Pagu", f"Rp {metrics['pagu']:,.0f}"); sub_c1.metric("Realisasi", f"Rp {metrics['realisasi']:,.0f}")
                        sub_c2.metric("Sisa", f"Rp {metrics['sisa']:,.0f}"); sub_c2.metric("Persentase", f"{metrics['persen']:.2f}%")

            st.markdown("---")
            
            # ==========================================
            # FITUR BARU: PROYEKSI PENYERAPAN (BURN-RATE FORECASTING)
            # ==========================================
            st.markdown("### 🔮 Burn-Rate Forecasting")
            
            active_months_list = [b for b, val in monthly_totals.items() if val > 0]
            num_active_months = len(active_months_list)
            
            if num_active_months > 0:
                avg_monthly_burn = total_realisasi_incl_out / num_active_months
                projected_end_year = avg_monthly_burn * 12
                projected_forecast_pct = (projected_end_year / total_pagu_all * 100) if total_pagu_all > 0 else 0
                
                f_col1, f_col2, f_col3 = st.columns(3)
                f_col1.metric("📅 Bulan Aktif Terdeteksi", f"{num_active_months} Bulan")
                f_col2.metric("⚡ Rata-rata Burn Rate / Bulan", f"Rp {avg_monthly_burn:,.0f}")
                f_col3.metric("🎯 Proyeksi Akhir Tahun", f"Rp {projected_end_year:,.0f}", f"{projected_forecast_pct:.2f}%")
                
                if projected_forecast_pct > 100: st.warning("⚠️ **Peringatan Over-Absorbtion**: Berdasarkan kecepatan penyerapan saat ini (*burn rate*), proyeksi pengeluaran melebihi total pagu anggaran. Perlu penyesuaian strategi.")
                elif projected_forecast_pct < 85: st.info("💡 **Catatan Under-Absorbtion**: Proyeksi penyerapan akhir tahun berada di bawah 85%. Akselerasi kegiatan di sisa bulan diperlukan untuk menghindari penumpukan anggaran.")
                else: st.success("✅ **Status Penyerapan Seimbang**: Proyeksi penyerapan berjalan optimal sesuai target anggaran.")
            else:
                st.info("Data bulanan aktif belum mencukupi untuk melakukan kalkulasi Burn-Rate Forecasting.")

            st.markdown("---")
            
            # ==========================================
            # GRAFIK BULANAN
            # ==========================================
            st.markdown("### 📊 REALISASI Bulanan")
            
            realisasi_per_bulan = [df_preview[bulan].max() if (bulan in df_preview.columns and not df_preview[bulan].empty) else 0 for bulan in list_semua_bulan]
            realisasi_per_bulan = [val if pd.notna(val) else 0 for val in realisasi_per_bulan]

            df_monthly_chart = pd.DataFrame({"Bulan": list_semua_bulan, "Realisasi": realisasi_per_bulan})
            df_monthly_chart["Realisasi_Juta"] = df_monthly_chart["Realisasi"] / 1_000_000
            
            def format_rupiah(val): return f"Rp {val:,.0f}".replace(",", ".") if val > 0 else "Rp 0"

            fig_3d_bar = px.bar(df_monthly_chart, x="Bulan", y="Realisasi_Juta", text=df_monthly_chart["Realisasi"].apply(format_rupiah), title="Realisasi Anggaran per Bulan", color="Realisasi_Juta", color_continuous_scale="Tealgrn")
            fig_3d_bar.update_traces(textposition='outside', textfont_size=11, marker_line_color='rgb(8,48,107)', marker_line_width=1.5, opacity=0.9)
            fig_3d_bar.update_layout(plot_bgcolor="rgba(245,247,250,0.8)", paper_bgcolor="rgba(0,0,0,0)", font=dict(color="black", size=12), xaxis_title="Bulan", yaxis_title="Total Realisasi (Juta Rp)", height=480, uniformtext_minsize=8, uniformtext_mode='hide')
            st.plotly_chart(fig_3d_bar, use_container_width=True)

            # ==========================================
            # DASHBOARD UI: REALISASI & SISA SUB KOMPONEN
            # ==========================================
            st.markdown("---")
            st.markdown("### 📋 Realisasi Berdasarkan Kegiatan")
            
            chart_colors = ['#1f77b4', '#ff7f0e', '#2ca02c', '#d62728', '#9467bd', '#8c564b', '#e377c2']
            
            if sub_component_realisasi:
                komp_keys = list(sub_component_realisasi.keys())
                sub_cols = st.columns(2)
                
                for idx, komp_name in enumerate(komp_keys):
                    with sub_cols[idx % 2]:
                        st.markdown(f"**{komp_name}**")
                        sub_dict = sub_component_realisasi[komp_name]
                        
                        sub_dict_filtered = {k: float(v) for k, v in sub_dict.items() if pd.notna(v) and float(v) >= 0}
                                
                        if len(sub_dict_filtered) > 0:
                            sub_labels = sorted(list(sub_dict_filtered.keys())) 
                            sub_values = [sub_dict_filtered[l] for l in sub_labels]
                            total_komp_val = sum(sub_values)
                            
                            if total_komp_val > 0:
                                fig_donut = go.Figure(data=[go.Pie(
                                    labels=[f"Sub {l.split('-')[0].strip()}" for l in sub_labels],
                                    values=sub_values,
                                    hole=0.45,
                                    textinfo='percent+label',
                                    hoverinfo='none',
                                    textfont=dict(size=11, color='#000000'),
                                    marker=dict(colors=chart_colors[:len(sub_labels)], line=dict(color="#332D2D", width=2))
                                )])
                                fig_donut.update_layout(plot_bgcolor="rgba(0,0,0,0)", paper_bgcolor="rgba(0,0,0,0)", margin=dict(t=20, b=20, l=20, r=20), showlegend=False, height=280)
                                st.plotly_chart(fig_donut, use_container_width=True)
                            else:
                                st.info("Seluruh sub komponen belum memiliki realisasi (Rp 0).")
                            
                            head_c = st.columns([0.5, 4.5, 2.5, 2])
                            with head_c[0]: st.markdown("")
                            with head_c[1]: st.markdown("**Sub Komponen**")
                            with head_c[2]: st.markdown("**Realisasi (Rp)**")
                            with head_c[3]: st.markdown("**%**")
                            st.markdown("<hr style='margin: 4px 0px 8px 0px;'>", unsafe_allow_html=True)

                            for i, (label, val) in enumerate(zip(sub_labels, sub_values)):
                                color_hex = chart_colors[i % len(chart_colors)]
                                pct = (val / total_komp_val * 100) if total_komp_val > 0 else 0
                                
                                row_c = st.columns([0.5, 4.5, 2.5, 2])
                                with row_c[0]: st.markdown(f"<div style='width:14px; height:14px; background-color:{color_hex}; border-radius:3px; margin-top:5px;'></div>", unsafe_allow_html=True)
                                with row_c[1]: st.markdown(f"<span style='font-size:12px; color:#212529;'>{label}</span>", unsafe_allow_html=True)
                                with row_c[2]: st.markdown(f"<span style='font-size:12px; font-weight:500; color:#212529;'>Rp {val:,.0f}</span>", unsafe_allow_html=True)
                                with row_c[3]: st.markdown(f"<span style='font-size:12px; font-weight:600; color:#495057;'>{pct:.2f}%</span>", unsafe_allow_html=True)
                        else:
                            st.info("Data realisasi belum tersedia (Rp 0).")
            else:
                st.info("Data realisasi sub komponen belum tersedia.")

            st.markdown("---")
            st.markdown("### 📋 Rincian Sisa Anggaran Per Kegiatan")
            
            if sub_component_sisa:
                komp_keys_sisa = list(sub_component_sisa.keys())
                sisa_cols = st.columns(2)
                
                for idx, komp_name in enumerate(komp_keys_sisa):
                    with sisa_cols[idx % 2]:
                        st.markdown(f"**{komp_name}**")
                        sisa_dict = sub_component_sisa[komp_name]
                        
                        sisa_dict_filtered = {k: float(v) for k, v in sisa_dict.items() if pd.notna(v) and float(v) >= 0}
                        
                        if len(sisa_dict_filtered) > 0:
                            sub_labels_sisa = sorted(list(sisa_dict_filtered.keys())) 
                            sub_values_sisa = [sisa_dict_filtered[l] for l in sub_labels_sisa]
                            total_komp_sisa_val = sum(sub_values_sisa)
                            
                            colors_sisa = ['#DAA520', '#CD853F', '#D2691E', '#B8860B', '#8B4513', '#A0522D', '#D2B48C']
                            
                            if total_komp_sisa_val > 0:
                                fig_donut_sisa = go.Figure(data=[go.Pie(
                                    labels=[f"Sub {l.split('-')[0].strip()}" for l in sub_labels_sisa],
                                    values=sub_values_sisa,
                                    hole=0.45,
                                    textinfo='percent+label',
                                    hoverinfo='none',
                                    textfont=dict(size=11, color='#000000'),
                                    marker=dict(colors=colors_sisa[:len(sub_labels_sisa)], line=dict(color="#272525", width=2))
                                )])
                                fig_donut_sisa.update_layout(plot_bgcolor="rgba(0,0,0,0)", paper_bgcolor="rgba(0,0,0,0)", margin=dict(t=20, b=20, l=20, r=20), showlegend=False, height=280)
                                st.plotly_chart(fig_donut_sisa, use_container_width=True)
                            else:
                                st.info("Seluruh anggaran telah terealisasi (Sisa Rp 0).")
                                
                            head_s = st.columns([0.5, 4.5, 2.5, 2])
                            with head_s[0]: st.markdown("")
                            with head_s[1]: st.markdown("**Sub Komponen**")
                            with head_s[2]: st.markdown("**Sisa (Rp)**")
                            with head_s[3]: st.markdown("**%**")
                            st.markdown("<hr style='margin: 4px 0px 8px 0px;'>", unsafe_allow_html=True)

                            for i, (label_sisa, val_sisa) in enumerate(zip(sub_labels_sisa, sub_values_sisa)):
                                color_hex_sisa = colors_sisa[i % len(colors_sisa)]
                                pct_sisa = (val_sisa / total_komp_sisa_val * 100) if total_komp_sisa_val > 0 else 0
                                
                                row_s = st.columns([0.5, 4.5, 2.5, 2])
                                with row_s[0]: st.markdown(f"<div style='width:14px; height:14px; background-color:{color_hex_sisa}; border-radius:3px; margin-top:5px;'></div>", unsafe_allow_html=True)
                                with row_s[1]: st.markdown(f"<span style='font-size:12px; color:#212529;'>{label_sisa}</span>", unsafe_allow_html=True)
                                with row_s[2]: st.markdown(f"<span style='font-size:12px; font-weight:500; color:#212529;'>Rp {val_sisa:,.0f}</span>", unsafe_allow_html=True)
                                with row_s[3]: st.markdown(f"<span style='font-size:12px; font-weight:600; color:#495057;'>{pct_sisa:.2f}%</span>", unsafe_allow_html=True)
                        else:
                            st.info("Seluruh anggaran telah terealisasi (Sisa Rp 0).")

            # ==========================================
            # DOWNLOAD BUTTONS
            # ==========================================
            st.divider()
            st.subheader("📥 DOWNLOAD LAPORAN")
            
            dl_col1, dl_col2 = st.columns(2)
            with dl_col1:
                st.download_button(
                    label="⬇️ Download Laporan Excel (.xlsx)",
                    data=output_excel_bytes,
                    file_name="LAPORAN_RAB_LRA_LENGKAP_JAN_DES.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary",
                    use_container_width=True
                )
            with dl_col2:
                st.download_button(
                    label="⬇️ Download Presentasi PowerPoint (.pptx)",
                    data=ppt_output_bytes,
                    file_name="PRESENTASI_EKSEKUTIF_ANGGARAN.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    type="primary",
                    use_container_width=True
                )

        except Exception as e:
            st.error(f"⚠️ Terjadi kesalahan pada saat pemrosesan: {e}")
